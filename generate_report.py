import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_human_report():
    doc = docx.Document()

    # Поля страницы
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    COLOR_TITLE = RGBColor(90, 30, 180)     # Приятный фиолетовый
    COLOR_BODY = RGBColor(35, 40, 50)        # Естественный темный цвет текста
    COLOR_MUTED = RGBColor(100, 110, 125)    # Серый для пояснений

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_BODY
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # --- ТИТУЛЬНАЯ ШАПКА ---
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r1 = p_title.add_run("ОТЧЕТ О ВЫПОЛНЕНИИ ПРАКТИЧЕСКИХ ЗАДАНИЙ\n")
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = COLOR_MUTED

    r2 = p_title.add_run("Разработка на Solana: Rust, Anchor, PDA и SPL Токены\n(Задания с 3 по 7 неделю)")
    r2.font.bold = True
    r2.font.size = Pt(18)
    r2.font.color.rgb = COLOR_TITLE

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Дисциплина: Startup / Web3 • Выполнил: Nerdlin")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_MUTED

    # Таблица сведений
    tbl_meta = doc.add_table(rows=3, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_meta.columns[0].width = Inches(2.2)
    tbl_meta.columns[1].width = Inches(4.7)

    meta_items = [
        ("Студент:", "Nerdlin"),
        ("Охват работы:", "Недели 3, 4, 5, 6, 7 учебного плана"),
        ("Используемый стек:", "Rust, Borsh, solana-program, Anchor 0.30, TypeScript, @solana/web3.js, @solana/spl-token, Metaplex"),
    ]
    for i, (k, v) in enumerate(meta_items):
        row = tbl_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_background(c0, "F3F4F6")
        set_cell_background(c1, "FAFAFA")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

        p = c0.paragraphs[0]
        r = p.add_run(k)
        r.font.bold = True
        r.font.size = Pt(10)

        p = c1.paragraphs[0]
        r = p.add_run(v)
        r.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Краткое введение простыми словами
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.add_run("О чем эта работа: ").font.bold = True
    p_intro.add_run(
        "В этом отчете собрано описание пяти практических заданий курса. "
        "Цель — последовательно пройти путь создания смарт-контрактов для Solana: "
        "разобраться с низкоуровневой моделью аккаунтов и сериализацией на Rust, задеплоить первую нативную "
        "программу в Devnet, перейти на современный фреймворк Anchor, реализовать хранение данных через PDA "
        "и настроить выпуск и переводы собственных токенов стандарта SPL."
    )

    def add_chapter(title, num):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(f"{num}. {title}")
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = COLOR_TITLE

    def add_note_box(text, title="Как проверить"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.9)
        set_cell_background(c, "F5F3FF")
        set_cell_margins(c, 80, 80, 120, 120)
        p = c.paragraphs[0]
        rt = p.add_run(f"📌 {title}: ")
        rt.font.bold = True
        rt.font.color.rgb = COLOR_TITLE
        rx = p.add_run(text)
        rx.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- НЕДЕЛЯ 3 ---
    add_chapter("Неделя 3: Rust для Solana-разработчика (CLI-симулятор)", "1")
    p = doc.add_paragraph()
    p.add_run("Что нужно было сделать: ").font.bold = True
    p.add_run("Написать небольшую CLI-программу на Rust с использованием базовых конструкций, которые понадобятся при разработке контрактов в Solana.")

    p = doc.add_paragraph()
    p.add_run("Что я сделал и почему именно так:\n").font.bold = True
    p.add_run(
        "Вместо абстрактных примеров я написал эмулятор рантайма Solana (папка week3_rust_cli), "
        "который моделирует, как блокчейн обрабатывает транзакции и изменяет состояние аккаунтов. "
        "В коде применены ключевые вещи, без которых смарт-контракты не пишутся:\n"
    )

    w3_points = [
        ("Сериализация Borsh: ", "Solana не использует тяжелый JSON — все данные сериализуются в компактный бинарный формат Borsh. В проекте через derive макросы запаковываются и инструкции (ProgramInstruction), и структуры состояния (CounterState, UserProfileState)."),
        ("Модель аккаунтов (Account & Pubkey): ", "Описал структуру аккаунта: 32-байтный публичный ключ, баланс в лампортах, буфер данных и флаги signer (подписана ли транзакция владельцем) и writable (можно ли менять)."),
        ("Точка входа процессора: ", "Функция Processor::process в точности повторяет сигнатуру стандартного entrypoint Solana: принимает program_id, срез аккаунтов и сырые байты инструкции."),
        ("Borrow Checker и split_at_mut: ", "В Rust нельзя одновременно взять две мутабельные ссылки на элементы одного вектора. Чтобы списать деньги с одного аккаунта и зачислить на другой, используется accounts.split_at_mut(1) — ровно так, как это делается в реальных контрактах."),
        ("Защищенная арифметика: ", "Все вычисления балансов и счетчиков сделаны через checked_add и checked_sub с возвратом ошибки ProgramError::ArithmeticOverflow при переполнении.")
    ]
    for k, v in w3_points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(k).font.bold = True
        bp.add_run(v)

    add_note_box("В папке week3_rust_cli выполнить: cargo run -- demo. Программа наглядно выполнит сериализацию, инициализирует счетчик, сделает перевод лампортов и покажет перехват ошибок при нехватке денег или отсутствии подписи.")

    # --- НЕДЕЛЯ 4 ---
    add_chapter("Неделя 4: Первая Solana Program на Devnet", "2")
    p = doc.add_paragraph()
    p.add_run("Что нужно было сделать: ").font.bold = True
    p.add_run("Создать и задеплоить свою первую программу в тестовую сеть Solana Devnet.")

    p = doc.add_paragraph()
    p.add_run("Как устроено решение:\n").font.bold = True
    p.add_run(
        "В папке week4_solana_program лежит нативный контракт на Rust и клиентский скрипт на TypeScript:\n"
    )

    w4_points = [
        ("Контракт (src/lib.rs): ", "Написан на базе библиотеки solana-program. Содержит точку входа entrypoint!(process_instruction) и логирование событий в ончейн msg!()."),
        ("Проверка владельца (Owner check): ", "Обязательное правило безопасности Solana: программа обязательно проверяет account.owner == program_id. Если аккаунт принадлежит кому-то другому, программа отказывается его менять и возвращает ошибку."),
        ("Состояние счетчика (GreetingAccount): ", "В аккаунте хранится количество вызовов программы и кастомная строка с сообщением, которую можно передать в транзакции."),
        ("Клиент на @solana/web3.js (client/client.ts): ", "Клиент подключается к Devnet, запрашивает аирдроп тестовых SOL, рассчитывает стоимость освобождения от арендной платы (Rent Exemption), создает аккаунт данных и отправляет транзакцию на вызов контракта."),
        ("Деплой: ", "Программу проще всего задеплоить через веб-инструмент Solana Playground (beta.solpg.io) — достаточно вставить код из src/lib.rs, нажать Build и Deploy.")
    ]
    for k, v in w4_points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(k).font.bold = True
        bp.add_run(v)

    add_note_box("Для взаимодействия с задеплоенной программой перейти в week4_solana_program/client, выполнить npm install и npx ts-node client.ts <ID_ПРОГРАММЫ>. Клиент выведет прямую ссылку на транзакцию в Solana Explorer.")

    # --- НЕДЕЛЯ 5 ---
    add_chapter("Неделя 5: Фреймворк Anchor (Counter Program)", "3")
    p = doc.add_paragraph()
    p.add_run("Что нужно было сделать: ").font.bold = True
    p.add_run("Написать программу-счетчик с использованием фреймворка Anchor и протестировать работу с ней через клиент.")

    p = doc.add_paragraph()
    p.add_run("Зачем нужен Anchor и что реализовано:\n").font.bold = True
    p.add_run(
        "В нативном Rust разработчику приходится писать сотни строк рутинного кода для проверки аккаунтов и парсинга байтов. "
        "Anchor решает эту проблему декларативными макросами #[derive(Accounts)]. В проекте week5_anchor_counter реализовано:\n"
    )

    w5_points = [
        ("Инструкции: ", "initialize (создание счетчика и привязка authority), increment (увеличение на 1), decrement (уменьшение на 1) и set_value (установка конкретного числа владельцем)."),
        ("Защита от ухода в минус: ", "В инструкции decrement добавлено строгое условие require!(counter.count > 0, ErrorCode::Underflow). Если счетчик уже равен 0, транзакция прерывается с понятной ошибкой."),
        ("Проверка прав: ", "В set_value использован констрейнт has_one = authority @ ErrorCode::Unauthorized. Никто, кроме создателя этого счетчика, не сможет принудительно изменить его значение."),
        ("Тесты на Mocha/Chai (tests/counter.ts): ", "Написаны тесты под все сценарии, включая проверку того, что ошибка Underflow действительно выбрасывается при попытке уйти в минус."),
        ("Клиент (client/interact.ts): ", "Готовый скрипт для вызова инструкций через сгенерированный IDL на Devnet.")
    ]
    for k, v in w5_points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(k).font.bold = True
        bp.add_run(v)

    add_note_box("Тесты запускаются командой anchor test (или во вкладке Test в Solana Playground). Клиент запускается из папки client через npm start.")

    # --- НЕДЕЛЯ 6 ---
    add_chapter("Неделя 6: Аккаунты, PDA и децентрализованный профиль", "4")
    p = doc.add_paragraph()
    p.add_run("Что нужно было сделать: ").font.bold = True
    p.add_run("Создать децентрализованный on-chain профиль пользователя с использованием механизма PDA.")

    p = doc.add_paragraph()
    p.add_run("Что такое PDA и как устроен профиль:\n").font.bold = True
    p.add_run(
        "PDA (Program Derived Address) — это одна из главных фишек Solana. Это адрес, который детерминированно "
        "вычисляется из строковых сидов и ID программы: seeds = [b'user-profile', wallet_pubkey]. "
        "У PDA нет приватного ключа — подписывать за него транзакции может только смарт-контракт. "
        "В папке week6_pda_user_profile реализовано:\n"
    )

    w6_points = [
        ("Гарантия одного профиля: ", "Так как адрес PDA привязан к публичному ключу кошелька, один пользователь физически не может наплодить дублирующие профили — адрес всегда один и тот же."),
        ("Хранилище данных (UserProfile): ", "В аккаунте хранятся authority (владелец), username (до 32 символов), bio (до 128 символов) и avatar_url (до 128 символов)."),
        ("Контроль размера строк: ", "Размеры полей жестко валидируются в контракте через require!, чтобы не допустить раздувания размера аккаунта."),
        ("Удаление с возвратом депозита (close constraint): ", "В инструкции delete_profile использован констрейнт close = authority. Когда пользователь решает удалить профиль, аккаунт закрывается в блокчейне, а все замороженные за аренду SOL возвращаются обратно на его кошелек.")
    ]
    for k, v in w6_points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(k).font.bold = True
        bp.add_run(v)

    add_note_box("Тест tests/user_profile.ts проверяет создание профиля, обновление данных, валидацию длины и факт возврата средств при закрытии аккаунта.")

    # --- НЕДЕЛЯ 7 ---
    add_chapter("Неделя 7: SPL Tokens (Выпуск, Transfer, Метаданные)", "5")
    p = doc.add_paragraph()
    p.add_run("Что нужно было сделать: ").font.bold = True
    p.add_run("Создать собственный токен: выпуск, transfer, metadata и взаимодействие через клиент.")

    p = doc.add_paragraph()
    p.add_run("Как работают токены в Solana:\n").font.bold = True
    p.add_run(
        "В Ethereum под каждый новый токен деплоят отдельный контракт стандарта ERC-20. "
        "В Solana все устроено гораздо быстрее и экономнее: в сети уже работает системная программа SPL Token Program. "
        "Мы создаем Mint-аккаунт (параметры самого токена) и ATA (Associated Token Account) — персональные хранилища "
        "токена для каждого конкретного пользователя.\n\n"
        "В папке week7_spl_tokens подготовлен скрипт, автоматизирующий весь процесс в Devnet:\n"
    )

    w7_points = [
        ("1. Создание токена (01_create_mint.ts): ", "Создание Mint аккаунта с точностью 9 десятичных знаков (decimals = 9)."),
        ("2. Метаданные Metaplex (02_create_metadata.ts): ", "Привязка официальных метаданных по стандарту Metaplex Token Metadata V3: название 'Startup University Token', символ 'STARTUP' и ссылка на логотип."),
        ("3. Чеканка токенов (03_mint_tokens.ts): ", "Создание связанного токен-аккаунта (ATA) для кошелька создателя и чеканка 50 000 токенов (MintTo)."),
        ("4. Перевод другому владельцу (04_transfer_tokens.ts): ", "Генерация второго кошелька (Bob), создание его ATA и перевод 1 500 токенов с подтверждением транзакции."),
        ("5. Проверка балансов (05_check_balances.ts): ", "Чтение финального состояния токена из блокчейна: общий объем предложения (Total Supply) и остатки на кошельках.")
    ]
    for k, v in w7_points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(k).font.bold = True
        bp.add_run(v)

    add_note_box("Сквозной запуск всего цикла: перейти в week7_spl_tokens и выполнить npm start. Скрипт сам запросит аирдроп SOL, создаст токен, отчеканит монеты, переведет часть на второй адрес и выведет ссылки на транзакции в Solana Explorer.")

    # --- ИТОГИ И ТАБЛИЦА ---
    add_chapter("Сводная таблица по всем заданиям", "6")

    tbl_res = doc.add_table(rows=6, cols=4)
    tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_res.columns[0].width = Inches(0.9)
    tbl_res.columns[1].width = Inches(1.8)
    tbl_res.columns[2].width = Inches(3.0)
    tbl_res.columns[3].width = Inches(1.2)

    hdr = ["Неделя", "Тема", "Что разработано", "Готовность"]
    row_hdr = tbl_res.rows[0]
    for idx, text in enumerate(hdr):
        c = row_hdr.cells[idx]
        set_cell_background(c, "5A1EB4")
        set_cell_margins(c, 80, 80, 80, 80)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    data_summary = [
        ("Неделя 3", "Rust для Solana", "CLI эмулятор рантайма, Borsh, Account model", "Готово ✅"),
        ("Неделя 4", "Первая программа", "Нативный смарт-контракт, Devnet деплой, клиент web3.js", "Готово ✅"),
        ("Неделя 5", "Anchor Framework", "Counter Program, Underflow require, тесты Mocha/Chai", "Готово ✅"),
        ("Неделя 6", "PDA и архитектура", "Децентрализованный профиль, возврат рента через close", "Готово ✅"),
        ("Неделя 7", "SPL Токены", "Скрипты создания минта, Metaplex V3, чеканка, трансфер", "Готово ✅"),
    ]

    for r_i, row_data in enumerate(data_summary, start=1):
        row = tbl_res.rows[r_i]
        bg = "FFFFFF" if r_i % 2 == 1 else "F9FAFB"
        for c_i, val in enumerate(row_data):
            cell = row.cells[c_i]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_i == 3:
                r.font.bold = True
                r.font.color.rgb = RGBColor(22, 101, 52)

    # Сохраняем
    target_path = r"c:\MI\пары\Startup\Отчет_по_заданиям_Solana_3-7.docx"
    doc.save(target_path)
    print(f"Обновленный человечный отчет сохранен: {target_path}")

if __name__ == "__main__":
    create_human_report()
